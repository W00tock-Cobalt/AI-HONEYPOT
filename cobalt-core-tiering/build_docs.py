#!/usr/bin/env python3
"""
Generates the Cobalt Core deliverables as Word (.docx) documents:

  1. Cobalt_Core_Pentester_Tiering_Framework.docx
  2. AI_Specialty_Vetting_Lab_Guide.docx

Run:  python3 build_docs.py
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- palette
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
COBALT = RGBColor(0x1A, 0x4D, 0xB3)
ACCENT = RGBColor(0x16, 0x82, 0x9B)
GREY = RGBColor(0x53, 0x5A, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = "EAF0FA"
LIGHTER = "F4F7FC"
HEADER_FILL = "0B1F3A"
BAND_FILL = "1A4DB3"


# ---------------------------------------------------------------- helpers
def set_cell_bg(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcpr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        m.append(node)
    tcpr.append(m)


def style_base(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(0x22, 0x26, 0x2C)
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.12


def add_run(p, text, size=10.5, bold=False, italic=False, color=None, font="Calibri"):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def h1(doc, text, color=COBALT):
    p = doc.add_paragraph()
    p.space_before = Pt(14)
    add_run(p, text, size=16, bold=True, color=color)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    # underline rule
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "1A4DB3")
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def h2(doc, text, color=NAVY):
    p = doc.add_paragraph()
    add_run(p, text, size=13, bold=True, color=color)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    return p


def h3(doc, text, color=ACCENT):
    p = doc.add_paragraph()
    add_run(p, text, size=11.5, bold=True, color=color)
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    return p


def para(doc, text, size=10.5, italic=False, color=None, after=6):
    p = doc.add_paragraph()
    add_run(p, text, size=size, italic=italic, color=color)
    p.paragraph_format.space_after = Pt(after)
    return p


def bullet(doc, text, level=0, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    if bold_lead:
        add_run(p, bold_lead, bold=True)
        add_run(p, text)
    else:
        add_run(p, text)
    p.paragraph_format.space_after = Pt(2)
    return p


def numbered(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Number")
    if bold_lead:
        add_run(p, bold_lead, bold=True)
        add_run(p, text)
    else:
        add_run(p, text)
    p.paragraph_format.space_after = Pt(2)
    return p


def callout(doc, text, fill=LIGHT, color=NAVY):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    cell.paragraphs[0].text = ""
    add_run(cell.paragraphs[0], text, size=10.5, italic=True, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def make_table(doc, headers, rows, widths=None, header_fill=HEADER_FILL,
               band=True, font_size=9.5, header_size=9.5):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, htext in enumerate(headers):
        set_cell_bg(hdr[i], header_fill)
        set_cell_margins(hdr[i])
        hdr[i].paragraphs[0].text = ""
        add_run(hdr[i].paragraphs[0], htext, size=header_size, bold=True, color=WHITE)
    for ridx, row in enumerate(rows):
        cells = tbl.add_row().cells
        fill = LIGHTER if (band and ridx % 2 == 0) else None
        for cidx, val in enumerate(row):
            set_cell_margins(cells[cidx])
            if fill:
                set_cell_bg(cells[cidx], fill)
            cells[cidx].paragraphs[0].text = ""
            # allow multi-line values via \n
            parts = str(val).split("\n")
            for pi, part in enumerate(parts):
                pgr = cells[cidx].paragraphs[0] if pi == 0 else cells[cidx].add_paragraph()
                bold = cidx == 0
                add_run(pgr, part, size=font_size, bold=bold)
                pgr.paragraph_format.space_after = Pt(0)
    if widths:
        for col_idx, w in enumerate(widths):
            for cell in tbl.columns[col_idx].cells:
                cell.width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def cover(doc, title, subtitle, meta_lines):
    # spacer
    for _ in range(3):
        doc.add_paragraph()
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = band.rows[0].cells[0]
    set_cell_bg(c, BAND_FILL)
    set_cell_margins(c, top=300, bottom=300, left=240, right=240)
    c.paragraphs[0].text = ""
    add_run(c.paragraphs[0], "COBALT CORE", size=13, bold=True, color=WHITE)
    tp = c.add_paragraph()
    add_run(tp, title, size=26, bold=True, color=WHITE)
    sp = c.add_paragraph()
    add_run(sp, subtitle, size=13, italic=True, color=RGBColor(0xD7, 0xE2, 0xF6))
    doc.add_paragraph()
    for label, value in meta_lines:
        p = doc.add_paragraph()
        add_run(p, f"{label}:  ", bold=True, color=NAVY)
        add_run(p, value, color=GREY)
        p.paragraph_format.space_after = Pt(2)
    doc.add_page_break()


def page_break(doc):
    doc.add_page_break()


# ================================================================
# DOCUMENT 1 — TIERING FRAMEWORK
# ================================================================
def build_framework():
    doc = Document()
    style_base(doc)
    sec = doc.sections[0]
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)

    cover(
        doc,
        "Penetration Tester Tiering",
        "A Capability-Based Staffing & Compensation Framework for Cobalt Core",
        [
            ("Document", "Penetration Tester Tiering Framework"),
            ("Owner", "Governance Team / Security Leadership"),
            ("Review cadence", "Bi-Annual recalibration"),
            ("Status", "Working draft for Core rollout"),
        ],
    )

    # ---- Overview
    h1(doc, "1. Overview")
    para(doc,
         "Penetration Tester Tiering at Cobalt Core is a structured framework designed to "
         "categorize security testing personnel based on their level of technical skill, breadth "
         "of specialties, and professional experience. The system uses three primary tiers — "
         "Tester (Tier 1), Senior Tester (Tier 2), and Specialist / Expert (Tier 3, the tester "
         "\u201Chat\u201D) — to define escalating proficiency in areas such as manual exploitation, "
         "complex architectural understanding, and custom tool development.")
    para(doc,
         "It also incorporates specialties — such as Web Security, Cloud Security, and AI/LLM "
         "Security — ensuring that a tester's placement is granular and reflects verified expertise "
         "in specific domains. Placement is validated through a structured vetting process involving "
         "lab environments and technical reporting.")
    callout(doc,
            "Core concept: shift the pen tester compensation and engagement model from a "
            "traditional contractor relationship to one that actively recognizes and rewards "
            "advanced, verified expertise \u2014 building a premium brand that attracts top-tier talent.")

    # ---- Strategic rationale
    h1(doc, "2. Strategic Rationale for a New Core Model")
    h2(doc, "Investment in Non-Employees")
    para(doc,
         "The company is intentionally investing in the development and compensation of its non-FTE "
         "testers, acknowledging that paying a higher rate is directly tied to the tester's proven "
         "capabilities. This is positioned as a rate for expertise, not a traditional employee wage.")
    h2(doc, "Talent Acquisition and Brand Building")
    para(doc,
         "The goal is to establish the company's \u201CCore\u201D testing community as the market "
         "leader, creating a self-fulfilling reputation: \u201Clike attracts like \u2014 once the "
         "reputation is set, they will come.\u201D By setting the market rate for high-skilled pen "
         "testing, the company aims to naturally draw the best talent, replacing the current sales "
         "and credit model that complicates attracting talent with a higher per-hour rate.")

    # ---- Mechanism
    h1(doc, "3. Mechanism for Expertise & Tiering")
    h2(doc, "Formal Assessment for Expertise")
    para(doc,
         "To formalize the designation of an \u201Cexpert\u201D and justify increased compensation, a "
         "standardized, skill-based assessment is implemented. Passing this assessment is the pathway "
         "to being recognized and compensated at the higher, expert-level tier.")
    h2(doc, "Compensation & Skill Tiers")
    para(doc,
         "The system is tied to a new tiered pay structure (e.g., standard pay rising to $75/hour, "
         "with subsequent tiers earning $15\u2013$20 hourly deltas). Tiered compensation aligns with "
         "the complexity of the work and the tester's proven skill level (e.g., compliance, senior, "
         "and boutique/enterprise), moving away from paying for work primarily done by automated scanners.")
    make_table(
        doc,
        ["Tier", "Designation", "Indicative Rate", "Work Profile"],
        [
            ["Tier 1", "Tester", "$75 / hour (baseline)", "Compliance-grade and standard assessments"],
            ["Tier 2", "Senior Tester", "+$15\u2013$20 / hour delta", "Senior, complex manual testing & chaining"],
            ["Tier 3", "Specialist / Expert", "+$15\u2013$20 / hour delta (per tier)", "Boutique / enterprise, specialty-led engagements"],
        ],
        widths=[0.9, 1.8, 2.0, 3.0],
    )
    para(doc, "Indicative figures are estimates for forecasting; placement is always confirmed through vetting.",
         italic=True, color=GREY, size=9)

    h2(doc, "Tier Maintenance")
    bullet(doc, "Tier 2 testers must perform X tests meeting Tier 2 requirements within a year to maintain Tier 2 status.",
           bold_lead="Senior (Tier 2): ")
    bullet(doc, "Specialized / Tier 3 testers must perform X tests within their specialty to maintain specialist status.",
           bold_lead="Specialist (Tier 3): ")
    para(doc,
         "If a tester does not perform the required number of tests, or their tier status otherwise "
         "lapses, they must go through the vetting process again to regain the status.")
    h2(doc, "Re-Vetting Conditions")
    para(doc, "Testers must go through the skill-validation process again if any of the following occur:")
    bullet(doc, "Their Tier 2 or Tier 3 status is allowed to lapse (maintenance threshold not met).")
    bullet(doc, "A specialty classification is allowed to lapse.")
    bullet(doc, "The governance team re-scopes a specialty's requirements during recalibration.")

    # ---- General tier competencies
    page_break(doc)
    h1(doc, "4. General Tier Competencies")
    para(doc, "The following competency bands apply across all specialties and methodologies.")

    h2(doc, "Tester (Tier 1)")
    bullet(doc, "Proficient using automated industry-standard tools (scanners, CSPM, SAST) for initial "
                "security-posture analysis and identifying common misconfigurations across environments "
                "(cloud, web, mobile, code).", bold_lead="Foundational technical & tool proficiency: ")
    bullet(doc, "Can manually identify, verify, and exploit basic, high-impact vulnerabilities (e.g., basic "
                "SQLi, IDORs, XSS, foundational network/Active Directory attacks) and understands the "
                "underlying protocols and security principles for their specialty.",
           bold_lead="Basic manual exploitation & analysis: ")
    bullet(doc, "Basic understanding of key benchmarks (OWASP Top 10, CIS/NIST) and can generate clear "
                "technical reports with actionable remediation advice.",
           bold_lead="Reporting & compliance awareness: ")

    h2(doc, "Senior Tester (Tier 2)")
    bullet(doc, "Deep manual testing to bypass modern defenses (WAFs, EDR/AV, SSL pinning) and exploit "
                "complex, non-obvious vulnerabilities \u2014 business-logic flaws, chaining multiple bugs to "
                "increase impact, complex data injection, and advanced specialized protocols (e.g., GraphQL, gRPC).",
           bold_lead="Advanced manual exploitation & logic-flaw discovery: ")
    bullet(doc, "Thorough manual review of critical architectural components automated tools miss \u2014 auditing "
                "complex IAM policies for privilege escalation, reviewing Infrastructure-as-Code, manual "
                "\u201Ctrace-to-source\u201D code analysis, and third-party library security (SCA).",
           bold_lead="Complex system & code review: ")
    bullet(doc, "Exploiting environment-specific vulnerabilities for greater impact \u2014 lateral movement in a "
                "cloud VPC, escaping basic container environments, advanced AD analysis (GPO abuse), and "
                "manipulating complex workflows (RAG, Tool-Use) in emerging technologies like AI/LLMs.",
           bold_lead="Environment & architectural deep-dive: ")

    h2(doc, "Specialist / Expert (Tier 3)")
    bullet(doc, "Expert / architect-level knowledge across the full attack surface \u2014 reverse engineering "
                "proprietary protocols or obfuscated binaries, complex low-level analysis (model weights, ARM "
                "assembly), and actively developing custom tools, extensions, or methodologies to discover and "
                "exploit novel vulnerabilities (including zero-days).",
           bold_lead="Deep technical mastery & custom research: ")
    bullet(doc, "Comprehensive architectural understanding of entire systems (multi-cloud, microservices, cloud "
                "identity boundaries) \u2014 designing organizational \u201Cguardrails,\u201D establishing "
                "secure-by-default architectures, defining the team's methodology, and acting as final reviewer "
                "for complex remediation and technical advisory.",
           bold_lead="Architectural & strategic security leadership: ")
    bullet(doc, "For leadership-focused roles: breaking down the technical scope of an engagement, organizing and "
                "distributing workload, providing technical direction for complex tasks, and delivering "
                "comprehensive reports with advanced technical reasoning.",
           bold_lead="Complex engagement management: ")

    # ---- Specialties matrix
    page_break(doc)
    h1(doc, "5. Specialties & Methodologies")
    para(doc, "Tiering is applied per specialty. A tester may hold different tiers across different specialties.")
    make_table(
        doc,
        ["#", "Specialty", "Focus"],
        [
            ["1", "Web Security", "OWASP Top 10, browser mechanics, manual exploitation"],
            ["2", "CCR (Cloud Config Review)", "Cloud control planes, identity, serverless, IaC"],
            ["3", "Cloud Security (AWS/Azure/GCP)", "Misconfigs, IAM, serverless, containerization"],
            ["4", "API Security", "Headless layer \u2014 REST, GraphQL, gRPC"],
            ["5", "Network Security (NetPen)", "Infrastructure, protocols, Active Directory"],
            ["6", "SCR (Secure Code Review)", "Source-code vulnerabilities (SAST / manual)"],
            ["7", "Mobile Security (iOS/Android)", "Package/binary analysis, local data storage"],
            ["8", "AI/LLM Security", "Large Language Models & prompt-based attacks"],
            ["9", "AI/MCP Security", "Integration layer between LLMs and external data/tools"],
            ["10", "AI Agentic Security Testing", "Autonomous agents, tool-using LLMs, multi-agent systems"],
            ["11", "Lead / Case Management", "Technical direction, communications, relationship mgmt"],
        ],
        widths=[0.4, 2.6, 4.0],
    )
    callout(doc,
            "Vetting note: Vetting for each specialty follows a similar format \u2014 testing within a lab "
            "environment and submitting a written report showcasing the work performed. Some specialties also "
            "require a questionnaire.",
            fill=LIGHTER)
    callout(doc,
            "Cloud note: Cloud specialty vetting is per cloud environment. Being vetted for Azure does not "
            "imply the same tier for AWS. A tester can hold multiple cloud tiers at multiple levels.",
            fill=LIGHTER)

    # ---- Per specialty detail
    specialties = [
        ("Web Security", "OWASP Top 10, browser mechanics, and manual exploitation.",
         "Proficient with automated scanners (Burp Suite, ZAP). Can manually identify and exploit basic "
         "vulnerabilities like reflected XSS, SQLi, and IDORs. Understands HTTP/S protocols and browser security headers.",
         "Performs deep manual testing. Can bypass WAFs, exploit complex CSRF/SSRF, and handle business-logic flaws. "
         "Proficient in JavaScript for DOM-based analysis and can chain multiple low-impact bugs into a high-impact exploit.",
         "Expert in modern frameworks (React, Angular). Performs deep source-code reviews, develops custom Burp "
         "extensions, and discovers zero-day vulnerabilities. Often final reviewer for complex remediation advice.",
         "Transition from current HTB labs to OffSec custom web tracks."),

        ("CCR (Cloud Config Review \u2014 AWS/Azure/GCP)",
         "Auditing cloud control planes, identity, serverless code, and infrastructure-as-code (IaC).",
         "Uses automated CSPM tools to identify deviations from benchmarks like CIS or NIST. Verifies low-hanging "
         "fruit such as unencrypted disks, public buckets, and overly permissive security groups.",
         "Audits complex IAM and resource-based policies for privilege-escalation paths. Reviews IaC (Terraform, "
         "CloudFormation, Bicep) before deployment to catch CI/CD misconfigurations. Understands cross-account trust.",
         "Architect-level understanding of cloud identity boundaries (SCPs, permissions boundaries). Designs "
         "organizational guardrails and identifies subtle flaws in multi-tenant/hybrid-cloud architectures. Expert "
         "in auditing KMS, secret managers, and private links.",
         "(Cobalt) Intentionally vulnerable cloud config, questionnaire, and report."),

        ("Cloud Security (AWS/Azure/GCP)",
         "Misconfigurations, IAM, vulnerable serverless functions, and containerization.",
         "Identifies common misconfigurations (public S3 buckets, open security groups). Understands the shared "
         "responsibility model and can run automated CSPM tools.",
         "Can exploit IAM role assumptions, perform lateral movement within a VPC, and escape basic container "
         "environments (Docker/K8s). Knowledgeable in cloud-native logging (CloudWatch/Azure Monitor) to avoid detection.",
         "Architect-level knowledge. Exploits complex multi-cloud environments, performs advanced Kubernetes cluster "
         "takeovers, and builds automated attack-and-detect simulations. Expert in serverless security and CI/CD pipeline poisoning.",
         "Present (March 2026): use current Hack The Box Pro offering for cloud evaluations. Future: transition to "
         "OffSec custom cloud tracks once available."),

        ("API Security", "The \u201Cheadless\u201D layer of applications (REST, GraphQL, gRPC).",
         "Understands RESTful principles and can interact with APIs. Identifies Broken Object Level Authorization "
         "(BOLA) and basic rate-limiting issues using automated tools.",
         "Proficient in testing GraphQL (introspection exploits) and gRPC. Performs complex data injection and "
         "bypasses sophisticated authentication/JWT implementations. Understands mass assignment and excessive data exposure.",
         "Expert in API gateway security and microservices architecture. Performs manual fuzzing of proprietary "
         "protocols and identifies race conditions or logic flaws in asynchronous API calls.",
         "Transition from current HTB labs to OffSec custom API tracks."),

        ("Network Security (NetPen)", "Infrastructure, protocols, and Active Directory.",
         "Executes a comprehensive manual methodology \u2014 LLMNR/NBT-NS poisoning, SMB Relay, Kerberoasting. "
         "Manual post-exploitation and pivoting (SSH/Proxychains/Chisel), and identifies misconfigurations in common "
         "services (SNMP, IPMI, FTP). Produces high-quality technical reports with actionable remediation.",
         "Complex environment takeovers and modern defensive bypass. Advanced AD analysis for non-obvious paths "
         "(GPO abuse, ACL delegation, unconstrained delegation). Bypasses NAC, executes VLAN hopping, and manually "
         "exploits legacy/proprietary services. Deep host-level EDR/AV evasion.",
         "Protocol-level exploitation and custom tool development. Reverse-engineers proprietary network protocols, "
         "performs manual cryptographic attacks on weak implementations, and exploits niche infrastructure "
         "(Mainframes, SCADA/ICS, VoIP). Develops internal methodology and provides architectural remediation for global networks.",
         "Transition from current HTB labs to OffSec custom NetPen tracks."),

        ("SCR (Secure Code Review)", "Identifying vulnerabilities within application source code (SAST / manual).",
         "Proficient with automated SAST tools (Checkmarx, Snyk, SonarQube). Triages noise from true positives for "
         "common flaws like SQLi and XSS. Understands basic secure-coding patterns in at least one major language "
         "(Java, Python, or C#).",
         "Performs manual trace-to-source analysis, following user input (sources) to execution points (sinks). "
         "Identifies complex logic flaws and authorization bypasses automated tools miss. Reviews multiple languages "
         "and understands third-party library security (SCA).",
         "Expert in identifying unreachable-code vulnerabilities and complex race conditions. Builds custom SAST "
         "rules (e.g., Semgrep) for proprietary business-logic flaws. Assists developers in architecting "
         "secure-by-default patterns and fixing systemic weaknesses.",
         "(Cobalt) Provide testers with intentionally vulnerable source code and a questionnaire; they evaluate the "
         "code for findings and produce a report."),

        ("Mobile Security (iOS/Android)", "Package/binary analysis and local data storage.",
         "Performs static analysis on APKs/IPAs using automated tools (MobSF). Identifies insecure local storage, "
         "hardcoded keys, and excessive permissions.",
         "Proficient in dynamic analysis using Frida or Objection. Bypasses SSL pinning and root/jailbreak detection. "
         "Understands IPC vulnerabilities and deep-link exploitation.",
         "Expert in reverse engineering obfuscated binaries and custom encryption. Performs manual exploit "
         "development for mobile OS kernels and analyzes complex ARM assembly code.",
         "Transition from current HTB labs to (Cobalt) intentionally vulnerable mobile applications, questionnaire, and report."),

        ("AI/LLM Security", "Large Language Models and prompt-based attacks.",
         "Understands the OWASP Top 10 for LLMs. Performs basic prompt injection (jailbreaking) and identifies insecure output handling.",
         "Performs indirect prompt injection and exploits data poisoning. Understands training-data extraction "
         "(model inversion) and manipulation of RAG (Retrieval-Augmented Generation) workflows.",
         "Deep understanding of model weights, gradients, and adversarial machine learning. Performs sophisticated "
         "evasion attacks against AI-driven defenses and audits neural-network architecture for supply-chain vulnerabilities.",
         "(Cobalt) Intentionally vulnerable LLM implementation, questionnaire, and report."),

        ("AI/MCP Security (Model Context Protocol)", "The integration layer between LLMs and external data/tools.",
         "Understands how MCP connects LLMs to local/remote resources. Identifies basic permission issues where an "
         "LLM is granted too much access to a filesystem or database via a host.",
         "Exploits Tool-Use vulnerabilities \u2014 tricking the model into executing unintended commands on the host "
         "via the MCP. Identifies flaws in the handoff between the Model, the Server, and the Client.",
         "Expert in the security architecture of the MCP standard itself. Identifies vulnerabilities in custom-built "
         "MCP servers, exploits sandbox escapes, and mitigates complex man-in-the-middle attacks between the LLM and "
         "its connected data sources.",
         "(Cobalt) Intentionally vulnerable LLM lab, questionnaire, and report."),
    ]
    for name, focus, t1, t2, t3, vet in specialties:
        h2(doc, name)
        para(doc, focus, italic=True, color=GREY, after=3)
        p = doc.add_paragraph(); add_run(p, "Tier 1: ", bold=True, color=COBALT); add_run(p, t1)
        p.paragraph_format.space_after = Pt(2)
        p = doc.add_paragraph(); add_run(p, "Tier 2: ", bold=True, color=COBALT); add_run(p, t2)
        p.paragraph_format.space_after = Pt(2)
        p = doc.add_paragraph(); add_run(p, "Specialist: ", bold=True, color=COBALT); add_run(p, t3)
        p.paragraph_format.space_after = Pt(2)
        p = doc.add_paragraph(); add_run(p, "Vetting: ", bold=True, color=ACCENT); add_run(p, vet, italic=True)
        p.paragraph_format.space_after = Pt(6)

    # AI Agentic — has Tier 1/2/3 distinctly named
    h2(doc, "AI Agentic Security Testing")
    para(doc, "Autonomous AI agents, tool-using LLMs, and multi-agent systems that can take actions in the real world.",
         italic=True, color=GREY, after=3)
    p = doc.add_paragraph(); add_run(p, "Tier 1 \u2014 Foundational: ", bold=True, color=COBALT)
    add_run(p, "Understands the OWASP Top 10 for LLM Applications and agentic AI threat models. Performs basic prompt "
               "injection (jailbreaking) and goal-hijacking attacks. Identifies insecure output handling and "
               "over-permissive tool access. Tests for privilege escalation through agent tool chains. Understands "
               "agent architectures: ReAct, AutoGPT, LangChain Agents, CrewAI.")
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph(); add_run(p, "Tier 2 \u2014 Advanced: ", bold=True, color=COBALT)
    add_run(p, "Performs indirect prompt injection via external data sources (documents, emails, web pages). Exploits "
               "tool/function-calling vulnerabilities (SSRF via agent, file-system access, code execution). Tests "
               "multi-agent communication for injection/manipulation between agents. Identifies confused-deputy "
               "attacks. Exploits memory/context poisoning in persistent sessions. Tests RAG poisoning and "
               "knowledge-base manipulation. Understands guardrail bypass and safety-filter evasion.")
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph(); add_run(p, "Tier 3 \u2014 Specialist: ", bold=True, color=COBALT)
    add_run(p, "Deep understanding of agent orchestration frameworks and their attack surfaces. Audits autonomous "
               "decision loops for unsafe action sequences. Exploits inter-agent trust boundaries in multi-agent "
               "swarms. Performs goal-misalignment attacks. Tests resource exhaustion and infinite-loop "
               "vulnerabilities. Audits human-in-the-loop bypass mechanisms. Understands supply-chain attacks on "
               "agent plugins, tools, and model weights. Evaluates sandboxing/containment for agent code execution.")
    p.paragraph_format.space_after = Pt(6)

    h2(doc, "Lead / Case Management")
    para(doc, "Technical direction, communications, and relationship management.", italic=True, color=GREY, after=3)
    p = doc.add_paragraph(); add_run(p, "Specialist: ", bold=True, color=COBALT)
    add_run(p, "Breaks down the testing tasks necessary to fulfill the brief and scope of each engagement. Organizes "
               "tasks into coverage categories and distributes workload to the testing team. Understands the technical "
               "requirements of the test and assists team staff with direction for complex work tasks. Experienced in "
               "writing findings and reports and can provide technical reasoning for all reported findings.")
    p.paragraph_format.space_after = Pt(6)

    # ---- Qualifications
    page_break(doc)
    h1(doc, "6. Qualifications & Eligibility")
    h2(doc, "Senior (Tier 2) \u2014 Baseline Qualifications")
    bullet(doc, "4+ years of professional baseline / career experience.")
    bullet(doc, "Web application testing proficient \u2014 ~60% of assessments.")
    bullet(doc, "Additional assessment eligibility: Network, AI/LLM, Cloud.")
    bullet(doc, "Coordinators (15 PT, 4.3 tester score).")
    bullet(doc, "Certifications, CVEs.")
    h3(doc, "Community Engagement (required)")
    bullet(doc, "Consistently submitting peer and pentest feedback post-pentest.")
    bullet(doc, "Active in Core Slack, surveys, office hours, etc.")
    bullet(doc, "Interested in producing content.")

    h2(doc, "Lead Testers (\u2248180)")
    bullet(doc, "Requires research across the past 3\u20135 tests.")
    bullet(doc, "Determine, from log review, what technical competencies qualify a tester as senior.")
    bullet(doc, "Has at least 3 Cobalt specialties \u2014 e.g., SCR, DRA, AI/LLM, Cloud.")
    bullet(doc, "Data-driven, supported by the data team.")
    bullet(doc, "Has submitted custom findings; reporting is clean.")
    bullet(doc, "Data team reviews vulnerability types, severity, and frequency.")
    bullet(doc, "Great client and team communication.")
    bullet(doc, "Contributes technical content and talks; an ambassador of Cobalt Core.")
    bullet(doc, "Goes above and beyond \u2014 proactive communication, consistent positive feedback from clients and TPMs.")
    bullet(doc, "High peer-feedback score.")
    bullet(doc, "Must upload ALL logs.")
    bullet(doc, "4 years with the Core.")

    h2(doc, "Hiring \u2014 Specialist / Expert Profile")
    bullet(doc, "6\u20138 years of professional pentesting experience.")
    bullet(doc, "Completes at least 8 tests a year.")
    bullet(doc, "Possible exploit-bounty program participation.")
    bullet(doc, "Operates at max capacity with a queue (q-in / q-out).")
    h3(doc, "Community Engagement (required)")
    bullet(doc, "Submits peer and pentest feedback after every test.")
    bullet(doc, "Actively provides feedback to Cobalt on application improvements, surveys, attends office hours.")
    bullet(doc, "Has produced or can contribute content to Cobalt.")
    h3(doc, "Funding & Customer-Lens Questions (open)")
    bullet(doc, "What does this pipeline look like? Does a client ask for this? Who pays for it? How do we fund the role?")
    bullet(doc, "As a pentest partner we bring in an elite/expert \u2014 so WE need to bring them in.")
    bullet(doc, "Validate, through the customer lens, how these people are the best.")
    bullet(doc, "More On-Demand \u2014 brought in to \u201Cclean up.\u201D")
    bullet(doc, "True top-1% folks brought in to hit the specialty hard (skills matrix).")
    h3(doc, "External Validation (public knowledge)")
    bullet(doc, "Regular presenter at Tier 2's.")
    bullet(doc, "Books, papers, and blogs to validate.")
    bullet(doc, "Websites; has written tools.")
    bullet(doc, "Active participation in local or international infosec clubs and conferences.")
    bullet(doc, "Must be publicly referenceable \u2014 no pseudonym or alias.")
    bullet(doc, "Specializes in a clear niche (the \u201CGoldilocks\u201D fit): Exploit Dev, SQLi, WAF, EDR bypasses, etc.")

    # ---- Transition / process
    page_break(doc)
    h1(doc, "7. Process: Transition to the Tiering System")
    para(doc,
         "We adopt a dual-track approach: a high-speed, needs-based process for the initial influx to "
         "quickly meet demand, followed by a more structured and comprehensive system for long-term, "
         "sustainable growth and quality control.")
    h2(doc, "Wave 1 \u2014 High-Velocity Integration")
    para(doc,
         "For the initial cohort, standard comprehensive tiering requirements can be temporarily relaxed. "
         "The priority is speed and immediate operational readiness:")
    bullet(doc, "Develop a high-level algorithmic filter focused only on the most critical, non-negotiable skills "
                "and compliance checks (e.g., minimum years of experience, essential certifications like OSCP, "
                "relevant professional history).", bold_lead="Algorithmic filter: ")
    bullet(doc, "Quick and focused on immediate deployment capability rather than exhaustive review of soft skills "
                "or niche competencies.", bold_lead="Expedited vetting: ")
    bullet(doc, "Onboarded testers receive a tier reflecting their fast-tracked status, allowing immediate "
                "contribution while deferring full requirements for a permanent, higher tier.",
           bold_lead="Temporary tier placement: ")
    h2(doc, "Recurring & Long-Term Strategy")
    bullet(doc, "Future hires evaluated against the complete criteria for each tier (Junior, Mid, Senior, Lead): "
                "technical expertise, report-writing quality, communication, project management, and domain knowledge.",
           bold_lead="Meet all established requirements: ")
    bullet(doc, "A clear, defined pathway for professional growth and tier advancement, ensuring fairness and transparency.",
           bold_lead="Standardized progression: ")
    bullet(doc, "The tiering structure and evaluation metrics are reviewed and updated to align with the evolving "
                "threat landscape and business needs.", bold_lead="Regular recalibration: ")

    h2(doc, "Data-Driven Determination")
    bullet(doc, "Evaluate the tester's historical reports, findings, communication, and technical skills.")
    bullet(doc, "Review the last 3\u20135 engagements; the tester must complete at least 5 engagements a year.")
    bullet(doc, "Specialty vetting: initially via HackTheBox, transitioning to Cobalt custom labs and OffSec custom "
                "tracks. The Research team interviews and reviews the tester's logs for any promotion request.")
    bullet(doc, "Cadence: testers can only be evaluated for promotion biannually, per the review cadence.")
    h3(doc, "Evaluation Criteria")
    make_table(
        doc,
        ["Criterion", "What it measures"],
        [
            ["Exploit sophistication", "Depth and novelty of exploitation beyond automated findings"],
            ["Chaining capability", "Combining multiple flaws into higher-impact attack paths"],
            ["Exploit velocity", "Time to achieve meaningful, verified impact"],
            ["QA rework rate", "Frequency of corrections required during quality assurance"],
            ["Severity stability", "Consistency / accuracy of severity ratings (low adjustment)"],
            ["Reporting clarity", "Clear, actionable, well-reasoned technical reporting"],
            ["SLA discipline", "Adherence to timelines and engagement commitments"],
        ],
        widths=[2.2, 5.0],
    )

    h2(doc, "Initial Tiering (Hard Cut-Off)")
    bullet(doc, "All current testers move to Tier 1.")
    bullet(doc, "Testers with a history of Informational / Low findings remain Tier 1.")
    bullet(doc, "Testers with extensive Lead experience likely qualify for Tier 2 or above.")
    bullet(doc, "Testers with a history of High / Critical findings likely qualify for Tier 2 or above.")
    bullet(doc, "Testers who actively engage, write blog posts, and go above and beyond may qualify for Tier 3 specialties.")
    bullet(doc, "Qualified testers with significant contributions and a proven delivery track record may be automatically vetted into Tier 2.")
    bullet(doc, "Qualified testers with prerequisite technical skills are offered the option to complete an assessment to attain Tier 3 specialties.")

    # ---- Recalibration table
    page_break(doc)
    h1(doc, "8. Bi-Annual Tiering Structure Recalibration")
    p = doc.add_paragraph()
    add_run(p, "Recalibration owner: ", bold=True); add_run(p, "Governance Team / Security Leadership.   ")
    add_run(p, "Cadence: ", bold=True); add_run(p, "Bi-annual, aligned with the overall governance and review cycle.")
    make_table(
        doc,
        ["Phase", "Focus Area", "Activities & Deliverables", "Stakeholders"],
        [
            ["1", "Threat Landscape Review",
             "Analyze emerging attack techniques and new technologies (AI/LLM advances, new cloud architectures). "
             "Review CVEs and industry reports from the past 6 months.\nDeliverable: summary of new/elevated threats "
             "and their impact on specialties.",
             "Research Team, Specialist Testers"],
            ["2", "Business & Capability Alignment",
             "Review company strategy, new product lines, and evolving client needs (enterprise vs. boutique). "
             "Assess whether the specialty list is complete and relevant.\nDeliverable: proposed adjustments to "
             "specialty definitions and tier caps.",
             "Leadership, Continuous Staffing Integration Team"],
            ["3", "Evaluation Metrics & Rubric Audit",
             "Review effectiveness of current evaluation criteria (e.g., exploit sophistication, QA rework rate). "
             "Audit the vetting process (HackTheBox / Cobalt Labs) per specialty for rigor.\nDeliverable: updated "
             "vetting-track requirements and metric weightings.",
             "Governance Team, Data Team, Research Team"],
            ["4", "Documentation & Communication",
             "Finalize updates to the tiering structure and documentation. Prepare a communication plan for Core "
             "testers on changes to compensation, vetting, and progression.\nDeliverable: approved, updated tiering "
             "document and internal communication package.",
             "Governance Team, HR/Legal, TPMs"],
        ],
        widths=[0.5, 1.6, 4.0, 1.6],
        font_size=9,
    )

    # ---- Governance & roadmap
    h1(doc, "9. Governance, Roadmap & Outcomes")
    h2(doc, "Governance & Economic Discipline")
    bullet(doc, "Tier caps enforced.")
    bullet(doc, "Evaluate data first.")
    bullet(doc, "Forecasting for finance (estimation only \u2014 we still vet for tiers).")
    bullet(doc, "Fast-track process for high-level testers.")
    bullet(doc, "Bi-annual reviews.")
    bullet(doc, "Specialty validation required.")
    bullet(doc, "Compensation strictly tier-aligned.")
    para(doc, "This ensures controlled evolution without hierarchy distortion.", italic=True, color=GREY)

    h2(doc, "Continuous Platform Alignment")
    para(doc, "Continuous Pentesting requires:")
    bullet(doc, "Multi-specialty validation.")
    bullet(doc, "Cross-sprint exploit consistency.")
    bullet(doc, "Reduced QA variability.")
    para(doc, "Capability segmentation becomes a platform-grade staffing infrastructure for Continuous.", italic=True, color=GREY)

    h2(doc, "Implementation Roadmap (6 Months)")
    make_table(
        doc,
        ["Phase", "Milestone"],
        [
            ["Phase 1", "Governance & Rubric Finalization"],
            ["Phase 2", "Cloud / CCR Pilot"],
            ["Phase 3", "Web / API Backscoring"],
            ["Phase 4", "SCR + AI Specialty Validation"],
            ["Phase 5", "Compensation Activation"],
            ["Phase 6", "Continuous Staffing Integration"],
        ],
        widths=[1.2, 6.0],
    )

    h2(doc, "Success Metrics (Measured Quarterly)")
    for m in ["QA rework hours per engagement", "Severity adjustment frequency", "Exploit chaining rate",
              "Time to meaningful findings", "Enterprise CSAT", "Continuous renewal rates", "Top-performer retention"]:
        bullet(doc, m)

    h2(doc, "Strategic Outcome")
    para(doc, "This program establishes:")
    for m in ["Capability-based staffing infrastructure", "Specialty-validated operators",
              "Controlled compensation differentiation", "Enterprise-grade consistency",
              "Continuous-ready Core architecture"]:
        bullet(doc, m)
    para(doc, "This represents the next phase of Core maturity aligned to Cobalt's product trajectory.", italic=True, color=GREY)

    h2(doc, "Tier Rollout")
    numbered(doc, "Announce the tiers.")
    numbered(doc, "Notify testers which tiers they are in.")
    numbered(doc, "Provide a period to challenge the tier.")
    numbered(doc, "Apply all objective criteria.")

    out = os.path.join(HERE, "Cobalt_Core_Pentester_Tiering_Framework.docx")
    doc.save(out)
    return out


# ================================================================
# DOCUMENT 2 — AI SPECIALTY VETTING LAB GUIDE
# ================================================================

# Module catalog: name -> dict(level, hours, category, rating fields, overview, vets)
MODULES = {
    "JarvisAI": dict(
        level="200", hours=1.0, category="Prompt Injection & AI Agent Abuse",
        signal=3.5, surface=3, difficulty="Foundational \u2013 Intermediate", overall=7.0,
        overview=(
            "A lab built around an AI-powered code-generation and execution service. Candidates interrogate "
            "an application that turns natural-language requests into executable logic, then must reason about "
            "what happens when the boundary between \u201Cgenerated code\u201D and \u201Cexecuted code\u201D is "
            "not properly enforced. From an initial foothold the scenario opens into host enumeration and a "
            "privilege-escalation path that rewards methodical post-exploitation. A clean, accessible warm-up "
            "that confirms a tester treats an AI feature as a real attack surface rather than a chatbot."),
        vets="AI-driven code-execution abuse, recon, web exploitation, Linux privilege escalation, post-exploitation discipline."),

    "PromptShock": dict(
        level="200", hours=1.0, category="Prompt Injection & AI Agent Abuse",
        signal=4.0, surface=4, difficulty="Intermediate", overall=7.5,
        overview=(
            "An AI-powered database assistant sits in front of sensitive backend data. The candidate must "
            "establish whether natural-language input can be steered into unintended backend actions, then use "
            "that leverage to enumerate hidden structure and recover sensitive artefacts. The scenario continues "
            "from initial access into a local privilege-escalation stage on the host. Strong signal for whether a "
            "tester understands the risk of coupling a language model to a privileged data layer."),
        vets="Prompt injection against backend logic, subdomain/recon enumeration, data-layer abuse, credential discovery, container/host privilege escalation."),

    "AImagery": dict(
        level="200", hours=1.0, category="Prompt Injection & AI Agent Abuse",
        signal=4.0, surface=4, difficulty="Intermediate", overall=8.0,
        overview=(
            "A multimodal AI image-analysis platform that processes user-submitted images and prompts. The "
            "candidate explores how AI-generated output is rendered and consumed downstream \u2014 including by "
            "privileged reviewers \u2014 and must chain an AI manipulation into a second-order client-side impact. "
            "Further enumeration surfaces data-exposure issues and a route to host access, followed by a "
            "misconfiguration-based privilege escalation. Note: services take time to initialise (allow ~15 "
            "minutes before testing). Excellent for assessing creative bug-chaining across the AI/output boundary."),
        vets="Prompt injection driving stored client-side execution, multimodal/data-exposure analysis, access-control flaws, credential recovery, misconfiguration-based privilege escalation, chaining."),

    "OmniCore": dict(
        level="200 (enterprise scope)", hours=1.0, category="Prompt Injection & AI Agent Abuse",
        signal=5.0, surface=5, difficulty="Intermediate \u2013 Advanced", overall=9.0,
        overview=(
            "A modern enterprise AI platform where several distinct, modern weaknesses chain from external "
            "access to full infrastructure compromise. The candidate begins by mapping the organisation's "
            "interconnected AI ecosystem, abuses an identity/normalisation flaw to gain elevated access through a "
            "vulnerable single-sign-on implementation, then weaponises trusted-content poisoning to manipulate an "
            "over-privileged assistant into disclosing restricted operational secrets. From there the path pivots "
            "into internal AI development infrastructure and culminates in a host-level compromise via insecure "
            "handling of serialized data in a privileged validation service. The single best proxy in the catalog "
            "for a realistic, end-to-end AI assessment."),
        vets="Identity/SSO abuse, indirect prompt injection & trust-boundary analysis, over-privileged LLM exploitation, AI infrastructure enumeration, insecure deserialization, end-to-end chaining."),

    "The Operator": dict(
        level="400", hours=4.0, category="Prompt Injection & AI Agent Abuse (Challenge)",
        signal=5.0, surface=5, difficulty="Advanced \u2013 Expert", overall=9.5,
        overview=(
            "A challenge-format scenario centred on an autonomous AI agent that researches third-party vendors "
            "and acts through tool calls (MCP). An attacker stands up a convincing but malicious vendor presence "
            "containing hidden instructions; when the agent ingests that external content, its task context is "
            "overridden and it silently performs unauthorised actions before resuming its normal workflow and "
            "producing a clean-looking report. The candidate must understand both how the manipulation is "
            "constructed and \u2014 critically \u2014 how to detect it, since the only reliable evidence lives in the "
            "raw tool-call telemetry rather than the agent's final output. Spans offence, detection, and incident "
            "response, making it an ideal capstone."),
        vets="Indirect prompt injection via external data, agentic/MCP tool-use abuse, autonomous-agent trust boundaries, exfiltration analysis, tool-call log forensics, detection-rule reasoning, IR."),

    "Megacorp One AI": dict(
        level="300", hours=2.0, category="Full AI Enterprise Environment",
        signal=4.5, surface=4, difficulty="Advanced", overall=8.5,
        overview=(
            "A full enterprise environment with AI/LLM components embedded in a broader, realistic corporate "
            "estate. Rather than a single isolated flaw, the candidate must navigate breadth \u2014 enumerating a "
            "wider attack surface, prioritising leads, and combining AI-specific weaknesses with conventional "
            "infrastructure and application findings to build impact. Strong for assessing whether a tester can "
            "sustain a methodical, multi-host engagement and integrate AI findings into a larger compromise narrative."),
        vets="Breadth-first enumeration, lead prioritisation, AI + conventional finding integration, multi-host engagement stamina, enterprise reporting."),

    "Graphflow": dict(
        level="200", hours=1.0, category="AI Agent Frameworks & MCP",
        signal=4.0, surface=4, difficulty="Intermediate", overall=8.0,
        overview=(
            "An MCP-powered natural-language-to-SQL workflow wired into a CI/CD context. The candidate examines "
            "how a model-driven query interface and its connected tooling handle untrusted input, and how an "
            "automation/delivery pipeline can amplify the impact of a model-layer weakness. A focused, modern "
            "look at the MCP integration layer and the blast radius created when agents touch build systems."),
        vets="MCP tool-use security, NL-to-SQL injection reasoning, CI/CD blast-radius analysis, integration-layer trust boundaries."),

    "MCP Inspector (CVE-2025-49596)": dict(
        level="200", hours=1.0, category="AI Agent Frameworks & MCP",
        signal=4.0, surface=4, difficulty="Intermediate", overall=8.0,
        overview=(
            "A scenario based on an unauthenticated remote-code-execution exposure in a widely used MCP tooling "
            "component. The candidate must recognise an exposed development/inspection surface in the MCP "
            "ecosystem, confirm the lack of authentication, and reason about how that exposure translates into "
            "host command execution. Tests whether a tester can move from \u201Cthis MCP tooling is reachable\u201D "
            "to demonstrable, responsible impact."),
        vets="MCP ecosystem attack surface, unauthenticated-service recognition, RCE confirmation and impact demonstration, responsible verification."),

    "BentoML Pickle RCE (CVE-2025-27520)": dict(
        level="200", hours=1.0, category="ML Framework & Inference API Exploitation",
        signal=4.0, surface=4, difficulty="Intermediate", overall=7.5,
        overview=(
            "A model-serving framework that mishandles serialized data on an inference pathway. The candidate "
            "must identify where untrusted input reaches an unsafe deserialization sink and reason about how that "
            "yields code execution on the serving host \u2014 a recurring, high-impact pattern across the ML "
            "infrastructure ecosystem. Good signal for whether a tester understands ML-serving internals, not just "
            "prompt-layer attacks."),
        vets="Insecure deserialization (pickle) recognition, ML inference-API internals, source-to-sink reasoning, serving-host RCE."),

    "LangFlow Unauthenticated RCE (CVE-2025-3248)": dict(
        level="200", hours=1.0, category="AI Platform CVEs",
        signal=4.0, surface=3.5, difficulty="Intermediate", overall=7.5,
        overview=(
            "A popular low-code AI application-builder platform that exposes a code-execution path without "
            "authentication. The candidate enumerates the platform, identifies the unauthenticated entry point, "
            "and reasons through to reliable host execution. Representative of the broad class of \u201CAI platform\u201D "
            "products that ship powerful execution features with weak access control."),
        vets="AI platform enumeration, unauthenticated-RCE identification, execution-path reasoning, AI-tooling supply-surface awareness."),

    "AgentScope eval() Abuse (CVE-2024-48050)": dict(
        level="200", hours=1.0, category="AI Agent Frameworks & MCP",
        signal=3.5, surface=4, difficulty="Intermediate", overall=7.0,
        overview=(
            "An agent framework that routes attacker-influenced input into a dynamic-evaluation sink. The "
            "candidate must trace how agent input reaches an unsafe evaluation primitive and turn that into "
            "command execution. A clean illustration of how \u201Ctool-use\u201D and dynamic evaluation create "
            "execution paths inside agent frameworks."),
        vets="Agent-framework code-injection, dynamic-eval sink analysis, tool-use abuse, source-to-sink tracing."),

    "DocsGPT MCP STDIO Abuse (CVE-2026-26015)": dict(
        level="200", hours=1.0, category="AI Agent Frameworks & MCP",
        signal=4.0, surface=4, difficulty="Intermediate", overall=7.5,
        overview=(
            "A document-QA assistant integrated over MCP via a standard-IO transport. The candidate examines the "
            "model-to-server-to-host handoff and how a transport/tool boundary can be abused to drive unintended "
            "behaviour on the host. A current example of the subtle trust issues introduced by MCP transports and "
            "tool wiring."),
        vets="MCP transport (STDIO) security, model/server/client handoff analysis, tool-boundary abuse, host-impact reasoning."),
}


def stars(score):
    full = int(round(score))
    full = max(0, min(5, full))
    return "\u2605" * full + "\u2606" * (5 - full)


def module_overview_block(doc, name):
    m = MODULES[name]
    h2(doc, name)
    p = doc.add_paragraph()
    add_run(p, "Level: ", bold=True, color=ACCENT); add_run(p, f"{m['level']}    ")
    add_run(p, "Duration: ", bold=True, color=ACCENT); add_run(p, f"{m['hours']:.1f} hr    ")
    add_run(p, "Category: ", bold=True, color=ACCENT); add_run(p, m["category"])
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    add_run(p, "Difficulty: ", bold=True, color=ACCENT); add_run(p, f"{m['difficulty']}    ")
    add_run(p, "Assessment signal: ", bold=True, color=ACCENT); add_run(p, f"{stars(m['signal'])}    ")
    add_run(p, "AI-surface focus: ", bold=True, color=ACCENT); add_run(p, f"{stars(m['surface'])}    ")
    add_run(p, "Overall: ", bold=True, color=ACCENT); add_run(p, f"{m['overall']:.1f}/10")
    p.paragraph_format.space_after = Pt(3)
    para(doc, m["overview"], after=3)
    p = doc.add_paragraph()
    add_run(p, "What it vets: ", bold=True, color=COBALT); add_run(p, m["vets"], italic=True)
    p.paragraph_format.space_after = Pt(8)


def track_table(doc, rows):
    total = sum(r[2] for r in rows)
    table_rows = [[r[0], r[1], f"{r[2]:.1f} hr", r[3]] for r in rows]
    table_rows.append(["", "TOTAL", f"{total:.1f} hr", ""])
    make_table(
        doc,
        ["#", "Module", "Time", "Primary competency assessed"],
        table_rows,
        widths=[0.4, 2.7, 0.8, 3.3],
        font_size=9,
    )


def build_ai_guide():
    doc = Document()
    style_base(doc)
    sec = doc.sections[0]
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)

    cover(
        doc,
        "AI Specialty Vetting \u2014 Lab Path Guide",
        "Module ratings & selectable assessment tracks for the AI/LLM, AI/MCP and AI Agentic specialties",
        [
            ("Document", "AI Specialty Vetting Lab Guide"),
            ("Companion to", "Cobalt Core Penetration Tester Tiering Framework"),
            ("Specialties served", "AI/LLM Security, AI/MCP Security, AI Agentic Security Testing"),
            ("Purpose", "Vet a tester's ability to perform a real AI security assessment"),
        ],
    )

    # ---- purpose
    h1(doc, "1. Purpose & How to Use This Guide")
    para(doc,
         "This guide supports the AI specialties in the Cobalt Core Penetration Tester Tiering Framework. It "
         "provides (a) a rating of every candidate lab module, (b) spoiler-free overviews of what each module "
         "assesses, and (c) four selectable assessment tracks \u2014 two at 8 hours and two at 10 hours \u2014 so "
         "the Research team can choose the format that best fits a given vetting cycle.")
    para(doc,
         "Consistent with the framework's vetting model, a candidate completes the chosen lab track and submits "
         "a written report (plus a questionnaire where required). Placement is confirmed by Research-team review "
         "of the report and the candidate's tool/MCP logs.")
    callout(doc,
            "Spoiler policy: Module write-ups describe the scenario, the attack surface, and the competencies "
            "assessed \u2014 but deliberately omit specific vulnerabilities, exploitation steps, and solution "
            "paths, so the labs remain valid as an assessment instrument.")

    # ---- scoring methodology
    h1(doc, "2. Rating Methodology")
    bullet(doc, "How strongly performance predicts real-world AI-assessment capability (1\u20135 stars).",
           bold_lead="Assessment signal: ")
    bullet(doc, "How much the module centres on genuinely AI-specific weaknesses versus generic web/host skills (1\u20135 stars).",
           bold_lead="AI-surface focus: ")
    bullet(doc, "Relative challenge, expressed as Foundational / Intermediate / Advanced / Expert.",
           bold_lead="Difficulty: ")
    bullet(doc, "A blended 0\u201310 score combining signal, AI-surface focus, realism, and chaining depth.",
           bold_lead="Overall: ")

    # ---- ratings summary table
    h1(doc, "3. Module Ratings \u2014 Summary")
    rows = []
    order = ["The Operator", "OmniCore", "Megacorp One AI", "AImagery", "Graphflow",
             "MCP Inspector (CVE-2025-49596)", "PromptShock", "BentoML Pickle RCE (CVE-2025-27520)",
             "LangFlow Unauthenticated RCE (CVE-2025-3248)", "DocsGPT MCP STDIO Abuse (CVE-2026-26015)",
             "JarvisAI", "AgentScope eval() Abuse (CVE-2024-48050)"]
    for name in order:
        m = MODULES[name]
        rows.append([name, m["level"], f"{m['hours']:.1f}h", stars(m["signal"]), stars(m["surface"]),
                     f"{m['overall']:.1f}"])
    make_table(
        doc,
        ["Module", "Level", "Time", "Signal", "AI-focus", "Overall"],
        rows,
        widths=[2.9, 1.0, 0.6, 1.0, 1.0, 0.7],
        font_size=8.5,
    )
    para(doc,
         "The five scenario labs (The Operator, OmniCore, AImagery, PromptShock, JarvisAI) form the strongest "
         "narrative core; Megacorp One AI adds enterprise breadth; the CVE-based labs add focused depth in MCP, "
         "ML-framework deserialization, and AI-platform exposures.", italic=True, color=GREY, size=9)

    # ---- mapping to tiering
    h1(doc, "4. Mapping to the Tiering Specialties")
    make_table(
        doc,
        ["Specialty", "Tier focus", "Representative modules"],
        [
            ["AI/LLM Security", "Tier 1\u20132: prompt injection, insecure output handling, indirect injection, RAG abuse",
             "JarvisAI, PromptShock, AImagery, OmniCore"],
            ["AI/MCP Security", "Tier 1\u20132: tool-use abuse, model/server/client handoff, MCP transport flaws",
             "Graphflow, MCP Inspector, DocsGPT MCP, (The Operator)"],
            ["AI Agentic Security", "Tier 2\u20133: indirect injection via external data, autonomous tool-use, detection",
             "The Operator, OmniCore, Megacorp One AI, AgentScope"],
        ],
        widths=[1.7, 3.0, 2.5],
        font_size=9,
    )
    para(doc,
         "Performance across a track should be read against the corresponding tier descriptors in the framework "
         "(e.g., a candidate who only completes direct-injection objectives demonstrates Tier 1 AI/LLM; a "
         "candidate who lands indirect injection, tool-use chaining, and detection demonstrates Tier 2\u20133).",
         italic=True, color=GREY, size=9)

    # ================= PART A: 8-HOUR =================
    page_break(doc)
    h1(doc, "PART A \u2014 8-Hour Assessment Tracks", color=NAVY)
    para(doc, "Two ways to assemble an 8-hour vetting cycle. Choose Track A1 for depth via fewer, larger scenarios, "
              "or Track A2 for breadth via eight one-hour modules.")

    h2(doc, "Track A1 \u2014 8 Hours, Multiple (Mixed-Length) Modules")
    para(doc, "Curated marquee scenarios of varying length. Best for assessing depth, realism, and end-to-end "
              "chaining \u2014 culminating in a four-hour agentic capstone.", italic=True, color=GREY)
    track_table(doc, [
        [1, "JarvisAI", 1.0, "AI code-execution abuse; warm-up"],
        [2, "PromptShock", 1.0, "Prompt injection into a privileged data layer"],
        [3, "AImagery", 1.0, "Multimodal injection \u2192 second-order client-side"],
        [4, "OmniCore", 1.0, "End-to-end enterprise AI compromise chain"],
        [5, "The Operator", 4.0, "Agentic indirect injection + detection (capstone)"],
    ])
    para(doc, "Coverage: direct & indirect prompt injection, multimodal, identity/SSO, insecure deserialization, "
              "agentic/MCP tool-use, and detection/IR.", size=9, color=GREY)

    h2(doc, "Track A2 \u2014 8 Hours, Exactly 8 Modules")
    para(doc, "Eight one-hour modules giving the widest coverage of the AI attack surface. Best for breadth and "
              "for sampling many distinct AI weakness classes quickly.", italic=True, color=GREY)
    track_table(doc, [
        [1, "JarvisAI", 1.0, "AI code-execution abuse"],
        [2, "PromptShock", 1.0, "Prompt injection \u2192 backend data"],
        [3, "AImagery", 1.0, "Multimodal injection \u2192 client-side"],
        [4, "OmniCore", 1.0, "Integrated enterprise AI chain"],
        [5, "Graphflow", 1.0, "MCP NL-to-SQL + CI/CD blast radius"],
        [6, "MCP Inspector (CVE-2025-49596)", 1.0, "Unauthenticated MCP-tooling RCE"],
        [7, "BentoML Pickle RCE (CVE-2025-27520)", 1.0, "ML-framework insecure deserialization"],
        [8, "LangFlow Unauth RCE (CVE-2025-3248)", 1.0, "AI-platform unauthenticated RCE"],
    ])
    para(doc, "Coverage: app-layer prompt injection, multimodal, MCP integration & tooling, ML-serving "
              "deserialization, and AI-platform exposures.", size=9, color=GREY)

    # ================= PART B: 10-HOUR =================
    page_break(doc)
    h1(doc, "PART B \u2014 10-Hour Assessment Tracks", color=NAVY)
    para(doc, "The same two formats extended to a 10-hour cycle for a more thorough vetting.")

    h2(doc, "Track B1 \u2014 10 Hours, Multiple (Mixed-Length) Modules")
    para(doc, "Track A1 plus a two-hour enterprise environment for added breadth before the agentic capstone.",
         italic=True, color=GREY)
    track_table(doc, [
        [1, "JarvisAI", 1.0, "AI code-execution abuse; warm-up"],
        [2, "PromptShock", 1.0, "Prompt injection into a privileged data layer"],
        [3, "AImagery", 1.0, "Multimodal injection \u2192 second-order client-side"],
        [4, "OmniCore", 1.0, "End-to-end enterprise AI compromise chain"],
        [5, "Megacorp One AI", 2.0, "Full enterprise environment; breadth + integration"],
        [6, "The Operator", 4.0, "Agentic indirect injection + detection (capstone)"],
    ])
    para(doc, "Coverage: everything in Track A1, plus sustained multi-host engagement stamina and integration of "
              "AI findings into a larger enterprise compromise.", size=9, color=GREY)

    h2(doc, "Track B2 \u2014 10 Hours, Exactly 10 Modules")
    para(doc, "Track A2 plus two more one-hour modules for deeper agent-framework and MCP coverage.",
         italic=True, color=GREY)
    track_table(doc, [
        [1, "JarvisAI", 1.0, "AI code-execution abuse"],
        [2, "PromptShock", 1.0, "Prompt injection \u2192 backend data"],
        [3, "AImagery", 1.0, "Multimodal injection \u2192 client-side"],
        [4, "OmniCore", 1.0, "Integrated enterprise AI chain"],
        [5, "Graphflow", 1.0, "MCP NL-to-SQL + CI/CD blast radius"],
        [6, "MCP Inspector (CVE-2025-49596)", 1.0, "Unauthenticated MCP-tooling RCE"],
        [7, "BentoML Pickle RCE (CVE-2025-27520)", 1.0, "ML-framework insecure deserialization"],
        [8, "LangFlow Unauth RCE (CVE-2025-3248)", 1.0, "AI-platform unauthenticated RCE"],
        [9, "AgentScope eval() Abuse (CVE-2024-48050)", 1.0, "Agent-framework dynamic-eval injection"],
        [10, "DocsGPT MCP STDIO (CVE-2026-26015)", 1.0, "MCP STDIO transport / handoff abuse"],
    ])
    para(doc, "Coverage: all of Track A2, plus agent-framework code injection and MCP transport-layer abuse.",
         size=9, color=GREY)

    callout(doc,
            "Choosing a track: Use the \u201Cmultiple modules\u201D tracks (A1 / B1) when you want to observe deep, "
            "realistic chaining and detection ability. Use the \u201C8 / 10 module\u201D tracks (A2 / B2) when you "
            "want maximum surface coverage and a finer-grained read on which AI weakness classes a candidate "
            "handles well.")

    # ================= MODULE OVERVIEWS =================
    page_break(doc)
    h1(doc, "5. Module Overviews (Spoiler-Free)")
    para(doc, "Every module referenced by any track, with its rating and a description of what the candidate will "
              "demonstrate \u2014 without revealing the solution.")

    h2(doc, "Scenario & Enterprise Labs")
    for name in ["JarvisAI", "PromptShock", "AImagery", "OmniCore", "Megacorp One AI", "The Operator"]:
        module_overview_block(doc, name)

    h2(doc, "MCP & Agent-Framework Labs")
    for name in ["Graphflow", "MCP Inspector (CVE-2025-49596)", "AgentScope eval() Abuse (CVE-2024-48050)",
                 "DocsGPT MCP STDIO Abuse (CVE-2026-26015)"]:
        module_overview_block(doc, name)

    h2(doc, "ML-Framework & AI-Platform Labs")
    for name in ["BentoML Pickle RCE (CVE-2025-27520)", "LangFlow Unauthenticated RCE (CVE-2025-3248)"]:
        module_overview_block(doc, name)

    # ================= ASSESSOR RUBRIC =================
    page_break(doc)
    h1(doc, "6. Assessor Rubric (Aligned to Framework Evaluation Criteria)")
    para(doc, "Score the candidate's report and logs against the same evaluation criteria used for tier "
              "determination, interpreted for AI assessments.")
    make_table(
        doc,
        ["Evaluation criterion", "What \u201Cstrong\u201D looks like in an AI assessment"],
        [
            ["Exploit sophistication", "Goes beyond a single jailbreak; demonstrates indirect injection, trust-boundary abuse, or deserialization with clear understanding of root cause."],
            ["Chaining capability", "Links AI-layer manipulation to concrete system impact (credential theft, RCE, host compromise) across multiple stages."],
            ["Exploit velocity", "Reaches meaningful, verified impact efficiently; does not stall on the AI feature as a novelty."],
            ["QA rework rate", "Findings are reproducible and accurate; minimal correction needed during review."],
            ["Severity stability", "Severities are justified and consistent with real business impact."],
            ["Reporting clarity", "Explains AI-specific root cause, attack path, evidence, and remediation in clear, actionable terms."],
            ["SLA discipline", "Completes the track within the allotted time and submits complete logs/report."],
            ["Detection & IR (agentic)", "For agentic modules, can locate evidence in tool-call telemetry and articulate a detection approach \u2014 not just the offence."],
        ],
        widths=[2.0, 5.2],
        font_size=9,
    )
    callout(doc,
            "Submission requirements (per framework): completed lab track + written report demonstrating the work; "
            "questionnaire where the specialty requires it; ALL logs uploaded for Research-team review.")

    out = os.path.join(HERE, "AI_Specialty_Vetting_Lab_Guide.docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    f1 = build_framework()
    f2 = build_ai_guide()
    print("Wrote:", f1)
    print("Wrote:", f2)
